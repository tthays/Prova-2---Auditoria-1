"""Gera planilhas editaveis (com formulas) dos papeis de trabalho quantitativos. Uso: python src/03_gerar_planilhas.py"""
import os, sys, json, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
sys.path.insert(0, os.path.dirname(__file__))
from config import *
RES = json.load(open(os.path.join(ROOT, "resultados", "resultados_resumo.json"), encoding="utf8"))
am = pd.read_csv(os.path.join(ROOT, "resultados", "amostra_mus_pt530.csv"))
reg = pd.read_csv(os.path.join(ROOT, "data", "registro_erros_inseridos.csv"))
F = "Arial"; AZ = Font(name=F, size=10, color="0000FF"); PR = Font(name=F, size=10); NG = Font(name=F, size=10, bold=True)
HD = PatternFill("solid", fgColor="1F3864"); HF = Font(name=F, size=10, bold=True, color="FFFFFF"); SUB = PatternFill("solid", fgColor="D9E2F3")
wb = Workbook()
def ws_new(nome, titulo, larg):
    ws = wb.create_sheet(nome); ws["A1"] = titulo; ws["A1"].font = Font(name=F, size=12, bold=True)
    ws["A2"] = f"Preparador: matricula {MATRICULA} | Data: {DATA_PREPARO} | Revisao: ____ (campo em branco, sem revisao de terceiros) | Azul = entrada/dado; preto = formula"; ws["A2"].font = Font(name=F, size=8, italic=True)
    for i, w in enumerate(larg, 1): ws.column_dimensions[get_column_letter(i)].width = w
    return ws
def hdr(ws, r, vals):
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v); c.font = HF; c.fill = HD; c.alignment = Alignment(wrap_text=True, vertical="center")
def put(ws, ref, v, fmt=None, azul=False, bold=False):
    c = ws[ref]; c.value = v; c.font = AZ if azul else (NG if bold else PR)
    if fmt: c.number_format = fmt
    return c
MIL = '#,##0;(#,##0);"-"'; PC = '0.0%'; RS = '#,##0.00;(#,##0.00);"-"'
wb.remove(wb.active)

# ================= PT 320 =================
w = ws_new("PT320", "PT 320 - Materialidade (R$ mil; DFP 2025 consolidada)", [46, 16, 16, 60])
hdr(w, 4, ["Base de referencia", "Valor 2025", "Percentual", "Fonte / justificativa"])
dados = [("Lucro antes dos tributos (LAIR)", DF["lair"][0], .05, "DFP p.23 - base escolhida: resultado recorrente e positivo; usuarios (acionistas Novo Mercado) focam lucro"),
         ("Receita liquida", DF["receita_liquida"][0], .01, "DFP p.23 - alternativa"), ("Ativo total", DF["ativo_total"][0], .005, "DFP p.20 - alternativa (empresa intensiva em ativos)"),
         ("Patrimonio liquido", DF["patrimonio_liquido"][0], .01, "DFP p.22 - alternativa")]
for i, (n_, v, p, f) in enumerate(dados, 5):
    put(w, f"A{i}", n_); put(w, f"B{i}", v, MIL, True); put(w, f"C{i}", p, PC, True); put(w, f"D{i}", f)
    put(w, f"E{i}", f"=B{i}*C{i}", MIL); 
hdr(w, 4, ["Base de referencia", "Valor 2025", "Percentual", "Fonte / justificativa", "Resultado (R$ mil)"]); w.column_dimensions["E"].width = 18
put(w, "A10", "Faixa das alternativas (minimo / maximo)", bold=True); put(w, "B10", "=MIN(E5:E8)", MIL); put(w, "C10", "=MAX(E5:E8)", MIL)
put(w, "A12", "MATERIALIDADE GLOBAL (5% do LAIR, arredondada)", bold=True); put(w, "B12", "=ROUND(E5,0)", MIL, bold=True)
put(w, "D12", "Dentro da faixa das alternativas (B10:C10) - teste de razoabilidade:"); put(w, "E12", '=IF(AND(B12>=B10,B12<=C10),"RAZOAVEL","REVISAR")')
put(w, "A13", "Percentual da materialidade para execucao (PM)"); put(w, "B13", .65, PC, True)
put(w, "D13", "65% (faixa usual 50-75%): auditoria inicial (PT 510), estimativa-chave de POC (PAA do relatorio, p.70) e partes relacionadas com risco elevado")
put(w, "A14", "MATERIALIDADE PARA EXECUCAO (PM)", bold=True); put(w, "B14", "=ROUND(B12*B13,0)", MIL, bold=True)
put(w, "A15", "Percentual do limite de distorcoes claramente triviais"); put(w, "B15", .05, PC, True)
put(w, "A16", "LIMITE DE DISTORCOES TRIVIAIS", bold=True); put(w, "B16", "=ROUND(B12*B15,0)", MIL, bold=True)
put(w, "A18", "Materialidades especificas (NBC TA 320 item 10)", bold=True)
put(w, "A19", "Partes relacionadas e remuneracao da administracao - % da global"); put(w, "B19", .10, PC, True)
put(w, "A20", "Materialidade especifica - partes relacionadas / honorarios"); put(w, "B20", "=ROUND(B12*B19,0)", MIL)
put(w, "D20", "Usuarios sao sensiveis a transacoes com partes relacionadas (nota 15, p.59) e a honorarios (nota 21, p.63)")
put(w, "A22", "Revisao: reavaliar B12 se o LAIR final divergir do usado em mais de 10% (NBC TA 320 itens 12-13). Nenhuma revisao ocorreu (didatico).")

# ================= PT 520 =================
w = ws_new("PT520", "PT 520 - Procedimentos analiticos (R$ mil)", [58, 16, 16, 16, 44])
put(w, "A4", "Limite de investigacao = 25% da PM (precisao da expectativa)"); put(w, "B4", "=0.25*PT320!B14", MIL)
put(w, "E4", "Diferencas acima do limite exigem investigacao")
hdr(w, 6, ["A. Encargos financeiros capitalizados em estoques (nota 7, p.51)", "2025", "2024", "", "Fonte"])
put(w, "A7", "Saldo inicial"); put(w, "B7", 111323, MIL, True); put(w, "C7", 80977, MIL, True)
put(w, "A8", "(+) Encargos capitalizados no exercicio"); put(w, "B8", 124065, MIL, True); put(w, "C8", 60088, MIL, True)
put(w, "A9", "(-) Baixa de encargos no exercicio"); put(w, "B9", -37935, MIL, True); put(w, "C9", -29742, MIL, True)
put(w, "A10", "Saldo final recalculado"); put(w, "B10", "=SUM(B7:B9)", MIL); put(w, "C10", "=SUM(C7:C9)", MIL)
put(w, "A11", "Saldo final divulgado"); put(w, "B11", 197453, MIL, True); put(w, "C11", 111323, MIL, True)
put(w, "A12", "Diferenca (recalculado - divulgado)"); put(w, "B12", "=B10-B11", MIL); put(w, "C12", "=C10-C11", MIL)
put(w, "E7", "DFP p.51, movimentacao dos encargos")
hdr(w, 14, ["B. Expectativa da baixa de encargos (proporcional ao saldo inicial)", "2025", "2024", "", ""])
put(w, "A15", "Razao baixa/saldo inicial em 2024"); put(w, "C15", "=-C9/C7", PC)
put(w, "A16", "Baixa esperada 2025 = saldo inicial 2025 x razao 2024"); put(w, "B16", "=B7*C15", MIL)
put(w, "A17", "Baixa registrada 2025"); put(w, "B17", "=-B9", MIL)
put(w, "A18", "Diferenca"); put(w, "B18", "=B17-B16", MIL); put(w, "A19", "Conclusao"); put(w, "B19", '=IF(ABS(B18)>$B$4,"INVESTIGAR","DENTRO DO LIMITE")')
hdr(w, 21, ["C. Taxa de capitalizacao de juros (notas 11 e 12, pp.55-56)", "2025", "2024", "", ""])
put(w, "A22", "Juros incorridos - emprestimos e financiamentos"); put(w, "B22", 81254, MIL, True); put(w, "C22", 57435, MIL, True)
put(w, "A23", "Juros incorridos - debentures e CRI"); put(w, "B23", 96628, MIL, True); put(w, "C23", 35994, MIL, True)
put(w, "A24", "Total de juros incorridos"); put(w, "B24", "=B22+B23", MIL); put(w, "C24", "=C22+C23", MIL)
put(w, "A25", "Encargos capitalizados no exercicio"); put(w, "B25", "=B8", MIL); put(w, "C25", "=C8", MIL)
put(w, "A26", "Razao capitalizado / incorrido"); put(w, "B26", "=B25/B24", PC); put(w, "C26", "=C25/C24", PC)
put(w, "A27", "Capitalizacao esperada 2025 = juros 2025 x razao 2024"); put(w, "B27", "=B24*C26", MIL)
put(w, "A28", "Diferenca (registrado - esperado)"); put(w, "B28", "=B25-B27", MIL); put(w, "A29", "Conclusao"); put(w, "B29", '=IF(ABS(B28)>$B$4,"INVESTIGAR","DENTRO DO LIMITE")')
put(w, "E28", "Investigacao exige dados internos (ativos qualificaveis por empreendimento): PLANEJADO")
hdr(w, 31, ["D. Indicadores (consolidado)", "2025", "2024", "2023", "Observacao"])
rows = [("Margem bruta", "=%d/%d" % (DF["lucro_bruto"][0], DF["receita_liquida"][0]), "=%d/%d" % (DF["lucro_bruto"][1], DF["receita_liquida"][1]), "=%d/%d" % (DF["lucro_bruto"][2], DF["receita_liquida"][2]), PC),
        ("Estoques totais (circ.+nao circ.)", DF["estoques_circ"][0] + DF["estoques_nc"][0], DF["estoques_circ"][1] + DF["estoques_nc"][1], DF["estoques_circ"][2] + DF["estoques_nc"][2], MIL),
        ("Custo dos imoveis vendidos", DF["custo"][0], DF["custo"][1], DF["custo"][2], MIL),
        ("Giro dos estoques em dias (estoque/CPV x 365)", "=B33/B34*365", "=C33/C34*365", "=D33/D34*365", '#,##0'),
        ("Contas a receber totais", DF["contas_receber_circ"][0] + DF["contas_receber_nc"][0], DF["contas_receber_circ"][1] + DF["contas_receber_nc"][1], DF["contas_receber_circ"][2] + DF["contas_receber_nc"][2], MIL),
        ("Receita liquida", DF["receita_liquida"][0], DF["receita_liquida"][1], DF["receita_liquida"][2], MIL),
        ("Prazo medio de recebimento em dias (CR/receita x 365)", "=B36/B37*365", "=C36/C37*365", "=D36/D37*365", '#,##0')]
for i, (n_, a, b_, c_, fm) in enumerate(rows, 32):
    put(w, f"A{i}", n_); [put(w, f"{col}{i}", v, fm, not str(v).startswith("=")) for col, v in zip("BCD", (a, b_, c_))]
put(w, "E32", "Margem +7,6 p.p. em 2025: INVESTIGAR (dados de empreendimento nao publicos)")
hdr(w, 40, ["E. Variacao de linhas selecionadas 2025 x 2024", "2025", "2024", "Var. %", "Investigar?"])
sel = [("Outras receitas operacionais (DRE, p.23)", 39857, -4134, "Notas 9(e),(f): ganhos AK14 R$16.641 e Tupi R$29.821 (liquidos) explicam parte - evidencia publica"),
       ("Despesas financeiras (DRE, p.23)", 86036, 44891, ""), ("Receitas financeiras (DRE, p.23)", 245765, 176919, ""),
       ("Contas a receber vencidas > 120 dias (nota 6, p.50)", 28016, 13130, ""), ("Provisao distratos e creditos (nota 6, p.50)", 24125, 11062, "")]
for i, (n_, a, b_, obs) in enumerate(sel, 41):
    put(w, f"A{i}", n_); put(w, f"B{i}", a, MIL, True); put(w, f"C{i}", b_, MIL, True); put(w, f"D{i}", f"=IF(C{i}=0,0,(B{i}-C{i})/ABS(C{i}))", PC)
    put(w, f"E{i}", f'=IF(ABS(B{i}-C{i})>$B$4,"INVESTIGAR","ok")&" {obs}"')

# ================= PT 530 =================
w = ws_new("PT530", "PT 530 - Amostragem MUS sobre a base SINTETICA (R$)", [44, 18, 18, 18, 18, 18, 22, 18])
m = RES["mus"]
hdr(w, 4, ["Parametro", "Valor", "Observacao"])
put(w, "A5", "Distorcao toleravel (= PM x 1.000, R$)"); put(w, "B5", "=PT320!B14*1000", RS)
put(w, "A6", "Fator de confianca (risco de aceitacao incorreta 5%, 0 erros)"); put(w, "B6", 3.00, '0.00', True)
put(w, "A7", "Intervalo de amostragem (R$)"); put(w, "B7", "=B5/B6", RS)
put(w, "A8", "Populacao: soma dos lancamentos positivos (R$)"); put(w, "B8", m["populacao_rs"], RS, True)
put(w, "A9", "Tamanho estimado da amostra"); put(w, "B9", "=ROUNDUP(B8/B7,0)", '0')
put(w, "A10", "Inicio aleatorio (semente 180109792)"); put(w, "B10", m["inicio_aleatorio"], RS, True)
put(w, "A11", "Itens-chave (valor >= intervalo)"); put(w, "B11", m["n_chave"], '0', True)
put(w, "A12", "Itens selecionados (script)"); put(w, "B12", "=COUNTA(A32:A600)", '0')
hdr(w, 14, ["Avaliacao (Stringer, Poisson 5%)", "taint", "Incremento", "Contribuicao (R$)"])
inc = [1.75, 1.55, 1.46, 1.40, 1.36, 1.33, 1.31, 1.29, 1.28, 1.27]
errs = am[(am.distorcao_rs > 0) & (~am.ja_identificada_teste_100pct)].sort_values("taint", ascending=False)
put(w, "A15", "Precisao basica = fator x intervalo"); put(w, "D15", "=B6*B7", RS)
for j, t in enumerate(errs.taint.values):
    r = 16 + j; put(w, f"A{r}", f"Excecao residual {j+1} (id {int(errs.id_lancamento.values[j])})"); put(w, f"B{r}", float(t), '0.0000', True)
    put(w, f"C{r}", inc[min(j, 9)], '0.00', True); put(w, f"D{r}", f"=B{r}*C{r}*$B$7", RS)
fim = 16 + len(errs)
put(w, f"A{fim}", "LIMITE SUPERIOR DA DISTORCAO (Stringer)", bold=True); put(w, f"D{fim}", f"=D15+SUM(D16:D{fim-1})", RS, bold=True)
put(w, f"A{fim+1}", "Projecao da distorcao (soma taint x intervalo)"); put(w, f"D{fim+1}", f"=SUMPRODUCT(B16:B{fim-1})*B7", RS)
put(w, f"A{fim+2}", "Limite superior x distorcao toleravel"); put(w, f"D{fim+2}", f'=IF(D{fim}>B5,"EXCEDE PM: ampliar amostra/ajustar","DENTRO DA PM")')
put(w, f"A{fim+3}", "Limite superior x materialidade global"); put(w, f"D{fim+3}", f'=IF(D{fim}>PT320!B12*1000,"EXCEDE GLOBAL","ABAIXO DA GLOBAL")')
hdr(w, 31, ["id_lancamento", "valor_rs", "valor_documento_simulado", "distorcao_rs", "tipo_excecao", "ja_identificada_100pct", "taint"])
for i, rr in enumerate(am.itertuples(), 32):
    for j, v in enumerate([rr.id_lancamento, rr.valor_rs, rr.valor_documento_simulado, rr.distorcao_rs, "" if pd.isna(rr.tipo_excecao) else rr.tipo_excecao, bool(rr.ja_identificada_teste_100pct), rr.taint], 1):
        c = w.cell(i, j, v); c.font = AZ
put(w, "I4", "Evidencia documental SIMULADA: valor_documento_simulado vem do registro de erros inseridos (oraculo), nao de documento real.")

# ================= PT 500 / conciliacao =================
w = ws_new("PT500", "PT 500 - Conciliacao da base sintetica com a DFP e verificacoes internas da DFP", [50, 20, 20, 20, 20, 22])
hdr(w, 4, ["Categoria (nota 7, p.51)", "Saldo contabil (R$)", "Base sintetica (R$)", "Diferenca (R$)", "Explicada pelo registro (R$)", "Nao explicada (R$)"])
cc = pd.read_csv(os.path.join(ROOT, "resultados", "conciliacao_por_categoria.csv"), index_col=0)
for i, (cat, rr) in enumerate(cc.drop("TOTAL").iterrows(), 5):
    put(w, f"A{i}", cat); put(w, f"B{i}", f"={NOTA7[cat]}*1000", RS); put(w, f"C{i}", float(rr.base_rs), RS, True); put(w, f"D{i}", f"=C{i}-B{i}", RS)
    put(w, f"E{i}", float(rr.explicada_pelo_registro_rs), RS, True); put(w, f"F{i}", f"=ROUND(D{i}-E{i},2)", RS)
put(w, "A10", "TOTAL", bold=True)
for col in "BCDEF": put(w, f"{col}10", f"=SUM({col}5:{col}9)", RS, bold=True)
hdr(w, 12, ["Verificacoes internas da DFP (real)", "Valor A", "Valor B", "Diferenca", "Resultado", "Fonte"])
chk = [("Ativo total = Passivo + PL (consolidado)", 7196637, "=396609+1680131+5119897", "p.20-22"),
       ("Ativo circulante + nao circulante = ativo total", "=3303420+3893217", 7196637, "p.20"),
       ("Nota 7 (soma) = estoques circ. + nao circ.", "=485897+1421820+1001151+197453+19513", "=1282510+1843324", "p.51 e p.20"),
       ("Nota 6 (contas a receber) = circ. + nao circ. no balanco", 1727447, "=439254+1288193", "p.50 e p.20"),
       ("Nota 12 debentures/CRI = circ. + nao circ.", 1097208, "=35558+1061650", "p.55 e p.21"),
       ("Nota 11 emprestimos = circ. + nao circ.", 608777, "=113237+495540", "p.55 e p.21"),
       ("PL: capital + reservas de capital + reservas de lucro + nao controladores", 5119897, "=4301053-62711+792204+89351", "p.22"),
       ("Caixa e equivalentes: balanco x DFC", 66575, 66574, "p.20 e p.27 (diferenca de arredondamento)"),
       ("Nota 7: encargos - roll forward", 197453, "=111323+124065-37935", "p.51"),
       ("Nota 11: emprestimos - movimentacao (836.309+429.214-670.620-67.380+81.254)", 608777, "=836309+429214-670620-67380+81254", "p.55"),
       ("Nota 12: debentures/CRI - movimentacao (300.073+757.911-59.306+96.628+1.902)", 1097208, "=300073+757911-59306+96628+1902", "p.55"),
       ("DFC captacoes (6.03.01) x notas 11+12 (429.214+757.911)", 1187114, "=429214+757911", "p.27, p.55 (dif. R$ 11 mil < trivial)"),
       ("Imobilizado+intangivel: saldo inicial + adicoes - deprec./amort. = saldo final", "=36152+2959+21775-12042", "=45743+3101", "p.20, p.26 (supoe ausencia de baixas)"),
       ("Nota 5: aplicacoes (CDB + fundos + LFs) = balanco", "=985389+40021+466557", 1491967, "p.49 e p.20"),
       ("Nota 21: remuneracao fixa + variavel = DRE (honorarios)", "=13835+6162", 19997, "p.63 e p.23"),
       ("Nota 14: tributos diferidos = circ. + nao circ.", 85425, "=18097+67328", "p.58")]
for i, (n_, a, b_, f) in enumerate(chk, 13):
    put(w, f"A{i}", n_); put(w, f"B{i}", a, MIL, not str(a).startswith("=")); put(w, f"C{i}", b_, MIL, not str(b_).startswith("="))
    put(w, f"D{i}", f"=B{i}-C{i}", MIL); put(w, f"E{i}", f'=IF(ABS(D{i})<=1,"CONFERE (<= R$ 1 mil)","DIVERGE")'); put(w, f"F{i}", f)

# ================= PT 450 =================
w = ws_new("PT450", "PT 450 - Mapa de distorcoes (SINTETICAS, inseridas) - avaliacao quantitativa (R$)", [34, 10, 22, 24, 40])
hdr(w, 4, ["Tipo de erro inserido", "Qtde", "Efeito na base (R$)", "Distorcao potencial (R$)", "Status (didatico)"])
mp = pd.read_csv(os.path.join(ROOT, "resultados", "mapa_distorcoes_pt450.csv"))
for i, rr in enumerate(mp.itertuples(), 5):
    put(w, f"A{i}", rr.tipo_erro); put(w, f"B{i}", int(rr.n), '0', True); put(w, f"C{i}", float(rr.efeito_na_base_rs), RS, True)
    put(w, f"D{i}", float(rr.distorcao_potencial_rs), RS, True); put(w, f"E{i}", "Nao corrigida (proposta de ajuste); detalhe em registro_erros_inseridos.csv")
fim = 5 + len(mp)
put(w, f"A{fim}", "TOTAL", bold=True)
for col in "BCD": put(w, f"{col}{fim}", f"=SUM({col}5:{col}{fim-1})", RS if col != "B" else '0', bold=True)
put(w, f"A{fim+2}", "Distorcao potencial agregada (R$ mil)"); put(w, f"D{fim+2}", f"=D{fim}/1000", MIL)
put(w, f"A{fim+3}", "Materialidade global (R$ mil)"); put(w, f"D{fim+3}", "=PT320!B12", MIL)
put(w, f"A{fim+4}", "Materialidade para execucao (R$ mil)"); put(w, f"D{fim+4}", "=PT320!B14", MIL)
put(w, f"A{fim+5}", "Limite de distorcoes triviais (R$ mil)"); put(w, f"D{fim+5}", "=PT320!B16", MIL)
put(w, f"A{fim+6}", "Agregado / global"); put(w, f"D{fim+6}", f"=D{fim+2}/D{fim+3}", PC)
put(w, f"A{fim+7}", "Conclusao quantitativa"); put(w, f"D{fim+7}", f'=IF(D{fim+2}>D{fim+3},"EXCEDE a materialidade global (se fossem reais: ajuste ou modificacao de opiniao)",IF(D{fim+2}>D{fim+4},"Excede a PM; abaixo da global: avaliar qualitativamente","Abaixo da PM"))')
put(w, f"A{fim+9}", "ATENCAO: distorcoes inseridas em dados sinteticos. Nao representam distorcoes da EZTEC.")

# ================= PT 540 =================
w = ws_new("PT540", "PT 540 - Sensibilidade da estimativa de custo orcado (POC) - nota 10, p.54 (R$ mil)", [56, 16, 16, 16, 16, 16])
put(w, "A4", "Receita de vendas contratadas (empreend. em construcao)"); put(w, "B4", 2889190, MIL, True)
put(w, "A5", "Custo orcado das unidades vendidas"); put(w, "B5", 1819076, MIL, True)
put(w, "A6", "Custo de construcao incorrido (sem encargos)"); put(w, "B6", 871239, MIL, True)
put(w, "A7", "POC (driver) = incorrido / orcado"); put(w, "B7", "=B6/B5", '0.00%'); put(w, "C7", "Divulgado: 47,89%")
put(w, "A8", "Receita apropriada (aproximacao = contratada x POC)"); put(w, "B8", "=B4*B7", MIL); put(w, "C8", "Divulgada (nota 10): 1.401.728")
hdr(w, 10, ["Cenario: variacao do custo orcado", "-10%", "-5%", "0%", "+5%", "+10%"])
for j, v in enumerate([-.10, -.05, 0, .05, .10]):
    col = get_column_letter(2 + j); put(w, f"{col}11", v, '0%', True)
    put(w, f"{col}12", f"=$B$5*(1+{col}11)", MIL); put(w, f"{col}13", f"=$B$6/{col}12", '0.00%'); put(w, f"{col}14", f"=$B$4*{col}13", MIL)
    put(w, f"{col}15", f"={col}14-$B$8", MIL); put(w, f"{col}16", f'=IF(ABS({col}15)>PT320!$B$14,"> PM","<= PM")')
for r_, t in zip(range(11, 17), ["Variacao do orcado", "Custo orcado ajustado", "POC ajustado", "Receita apropriada ajustada", "Efeito na receita/resultado bruto", "vs. PM"]): put(w, f"A{r_}", t)
put(w, "A18", "Variacao do orcado que consome a PM (break-even)"); put(w, "B18", "=B6/(B7-PT320!B14/B4)/B5-1", '0.00%')
put(w, "A19", "Limitacao: aproximacao agregada; o POC real e calculado por unidade/empreendimento (nota 2.16, p.46). Orcado por empreendimento nao e publico: PLANEJADO.")

# ================= PT 570 =================
w = ws_new("PT570", "PT 570 - Continuidade operacional: liquidez e endividamento (R$ mil; consolidado)", [62, 16, 16, 16, 50])
hdr(w, 4, ["Item", "2025", "2024", "2023", "Fonte"])
ln = [("Ativo circulante", "ativo_circulante"), ("Passivo circulante", "passivo_circulante"), ("Caixa e equivalentes", "caixa_equiv"), ("Aplicacoes financeiras", "aplicacoes_fin"),
      ("Emprestimos, financiamentos e debentures - circulante", "emprest_circ"), ("Emprestimos, financiamentos e debentures - nao circulante", "emprest_nc"),
      ("Patrimonio liquido", "patrimonio_liquido"), ("Caixa liquido das atividades operacionais", "fco")]
for i, (n_, k) in enumerate(ln, 5):
    put(w, f"A{i}", n_); [put(w, f"{c}{i}", DF[k][j], MIL, True) for j, c in enumerate("BCD")]; put(w, f"E{i}", "DFP " + DF[k][3])
put(w, "A14", "Liquidez corrente"); put(w, "A15", "Disponibilidades (caixa + aplicacoes)"); put(w, "A16", "Liquidez imediata (disponibilidades / passivo circulante)")
put(w, "A17", "Divida bruta"); put(w, "A18", "Divida liquida (bruta - disponibilidades)"); put(w, "A19", "Divida bruta / PL"); put(w, "A20", "Divida liquida / PL")
put(w, "A21", "Disponibilidades / divida circulante (cobertura do curto prazo)")
for c in "BCD":
    put(w, f"{c}14", f"={c}5/{c}6", '0.00'); put(w, f"{c}15", f"={c}7+{c}8", MIL); put(w, f"{c}16", f"={c}15/{c}6", '0.00'); put(w, f"{c}17", f"={c}9+{c}10", MIL)
    put(w, f"{c}18", f"={c}17-{c}15", MIL); put(w, f"{c}19", f"={c}17/{c}11", PC); put(w, f"{c}20", f"={c}18/{c}11", PC); put(w, f"{c}21", f"={c}15/{c}9", '0.0"x"')
hdr(w, 23, ["Covenant (nota 12, p.56): (Divida liquida + imoveis a pagar) / PL < 0,80", "Valor", "", "", "Observacao"])
put(w, "A24", "Debentures e CRI (liquido de custos)"); put(w, "B24", 1097208, MIL, True); put(w, "A25", "(-) Disponibilidades 2025"); put(w, "B25", "=-B15", MIL)
put(w, "A26", "(+) Imoveis a pagar (terrenos a pagar circ. 36.421 + nao circ. 28.027)"); put(w, "B26", "=36421+28027", MIL)
put(w, "A27", "Numerador recomputado (HIPOTESE: exclui financiamentos a producao)"); put(w, "B27", "=SUM(B24:B26)", MIL)
put(w, "A28", "Indice recomputado"); put(w, "B28", "=B27/B11", '0.00'); put(w, "E28", "Divulgado: (0,08). Conferencia aproximada; definicao contratual nao acessada")
put(w, "A29", "Segundo indice (>1,50): divulgado 13,83 - nao reproduzivel com dados publicos: PLANEJADO (escritura de emissao)")
put(w, "A31", "Fluxo: FCO positivo em 2025 apos dois exercicios negativos; captacoes liquidas financiaram investimentos (DFC p.26-27).")

# ================= PT 505 =================
w = ws_new("PT505", "PT 505 - Controle de circularizacao (DIDATICO: NENHUMA solicitacao enviada)", [34, 34, 18, 14, 22, 22, 36])
hdr(w, 4, ["Tipo", "Saldo a confirmar (DFP)", "Valor (R$ mil)", "Modelo", "Status", "Data envio", "Tratamento de diferencas / procedimento alternativo"])
cir = [("Bancos - aplicacoes financeiras", "Nota 5 (p.49)", 1491967, "Modelo A", "Aplicacoes e saldos em conta; garantias/gravames"),
       ("Bancos - caixa e equivalentes", "Nota 4 (p.49)", 66575, "Modelo A", "Extratos e conciliacoes bancarias"),
       ("Credores - emprestimos e financiamentos", "Nota 11 (p.55)", 608777, "Modelo A", "Contratos, extratos de credor, recalculo de juros"),
       ("Debentures/CRI - agente fiduciario", "Nota 12 (p.55)", 1097208, "Modelo A", "Extrato do agente fiduciario e escrituras"),
       ("Clientes - contas a receber", "Nota 6 (p.50)", 1727447, "Modelo B", "Positiva por amostra MUS; alternativos: contratos, recebimentos posteriores"),
       ("Advogados - litigios", "Nota 16 (p.60)", 46822, "Modelo C", "Carta de advogados: provaveis (6.469) e possiveis (46.822)"),
       ("Partes relacionadas - saldos", "Nota 15 (p.59)", 77302, "Modelo D", "Confirmacao com contraparte e contratos"),
       ("Terrenos a pagar", "BP (p.21-22)", 64448, "Modelo D", "Contratos de compra e recibos")]
for i, (t, s, v, mo, tr) in enumerate(cir, 5):
    put(w, f"A{i}", t); put(w, f"B{i}", s); put(w, f"C{i}", v, MIL, True); put(w, f"D{i}", mo); put(w, f"E{i}", "NAO ENVIADO (didatico)"); put(w, f"F{i}", "-"); put(w, f"G{i}", tr)
put(w, "A14", "Total do escopo planejado (R$ mil)"); put(w, "C14", "=SUM(C5:C12)", MIL)
put(w, "A15", "Obs.: total mistura saldos de ativo e passivo e exposicoes; nao representa cobertura percentual de uma demonstracao.")

wb.save(os.path.join(ROOT, "planilhas", "Planilhas_P2_180109791.xlsx")); print("ok")
